from io import BytesIO
from docx import Document
def make_docx(title: str, content: str) -> bytes:
    doc = Document(); doc.add_heading(title or "DraftMate Document", 0)
    for block in content.split("\n"):
        line = block.strip()
        if not line: continue
        if line.startswith("### "): doc.add_heading(line[4:], level=3)
        elif line.startswith("## "): doc.add_heading(line[3:], level=2)
        elif line.startswith("# "): doc.add_heading(line[2:], level=1)
        else: doc.add_paragraph(line)
    out=BytesIO(); doc.save(out); return out.getvalue()
