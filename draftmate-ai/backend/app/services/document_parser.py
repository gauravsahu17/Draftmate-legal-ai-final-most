from io import BytesIO
from pathlib import Path
from docx import Document as DocxDocument
from pypdf import PdfReader
ALLOWED = {".docx", ".pdf", ".txt"}
def extract_document(filename: str, data: bytes) -> tuple[str, str]:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED: raise ValueError("Unsupported file type. Upload DOCX, text-based PDF, or TXT.")
    if not data: raise ValueError("The uploaded file is empty.")
    if ext == ".txt": text = data.decode("utf-8-sig", errors="replace")
    elif ext == ".docx":
        doc = DocxDocument(BytesIO(data)); parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                style = p.style.name.lower() if p.style else ""
                parts.append(("# " if "heading 1" in style else "## " if "heading" in style else "") + p.text)
        for table in doc.tables: parts.extend(" | ".join(c.text.replace("\n", " ") for c in row.cells) for row in table.rows)
        text = "\n\n".join(parts)
    else:
        reader = PdfReader(BytesIO(data)); text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip(): raise ValueError("No selectable text was found. This may be a scanned PDF; OCR is not configured.")
    text = text.strip()
    if not text: raise ValueError("No readable text was found in this document.")
    return text, ext[1:].upper()
